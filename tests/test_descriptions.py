"""The description round-trip: a person writes meanings, the program joins them.

    python tests/test_descriptions.py

No API key, no quota, no network, and nothing stubbed — `descriptions.py` is
pure Python over records, so this exercises the real thing end to end.

Two properties are load-bearing, and both fail quietly if they break:

  1. **The overlay never enters the records.** `output.json` means exactly one
     thing: what the source documents state, every value checkable
     character-for-character against the one file it came from. A sentence a
     person wrote cannot be checked that way. Merging it would leave
     `validate_assembled` printing "N/N records clean" over a dictionary that
     is now part testimony, and turn `with_description` into a permanent 100%.
     So the join happens in the spreadsheet, and the JSON is untouched.

  2. **One description per FIELD, fanned out to every record.** One record is
     one (target field, source field) pair, so a field with ten sources is ten
     records — but it has one meaning. She writes it once.

The rest is hostility toward a spreadsheet round-trip, because the medium is
Excel and the author is a human: case, padding, non-breaking spaces, NFD vs
NFC, re-sorted rows, and keys that no longer exist.
"""

import json
import shutil
import sys
import tempfile
import unicodedata
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

import descriptions
from assemble import short_path
from descriptions import AUTHORED_COL, DOCUMENT_COL
from emit import write_workbook

# --- synthetic records, in exactly the shape assemble.py builds -------------
#
# WIDE is n->1: one target field fed by three source fields, so three records
# share one meaning. DESCRIBED already carries a document-stated description,
# so the two description columns have something to keep apart.

STAGES = ["source", "staging", "dwh"]


def _entry(stage, table, column, sources=(), path="/a/Archive/H/T.xlsx"):
    return {"stage": stage, "schema": None, "table": table, "column": column,
            "datatype": "VARCHAR2", "size": "20", "constraint": None,
            "doc_link": None, "offline_path": path, "sources": list(sources)}


def _source(table, column):
    return {"table": table, "column": column,
            "transformation_type": "Direct Mapping",
            "transformation_logic": "1-1", "role": "value"}


def _record(column, source_column, description=None):
    return {"description": description, "lineage": [
        _entry("source", "RAW_PARTY", source_column),
        _entry("dwh", "DWH_PARTY", column, [_source("RAW_PARTY", source_column)]),
    ]}


RECORDS = (
    # one field, three sources -> three records, one meaning
    [_record("WIDE_FIELD", f"SRC_{i}") for i in range(3)]
    + [_record("PLAIN_FIELD", "SRC_P")]
    + [_record("DESCRIBED_FIELD", "SRC_D", "Ngày mở tài khoản")]
)

WIDE = ("DWH_PARTY", "WIDE_FIELD")


def cell_of(ws, row, header):
    return ws.cell(row, [c.value for c in ws[1]].index(header) + 1)


def fill(path, edits, extra_rows=()):
    """Play the human: type into the authored column, then add stray rows."""
    wb = load_workbook(path)
    ws = wb["Descriptions"]
    headers = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        key = ((row[headers.index("Bảng")].value or "").strip().upper(),
               (row[headers.index("Cột")].value or "").strip().upper())
        if key in edits:
            table, column, text = edits[key]
            row[headers.index("Bảng")].value = table
            row[headers.index("Cột")].value = column
            row[headers.index(AUTHORED_COL)].value = text
    for extra in extra_rows:
        ws.append(extra)
    wb.save(path)


def main_():
    tmp = Path(tempfile.mkdtemp(prefix="hecate_desc_"))
    failed = 0
    try:
        template = tmp / "mota.xlsx"
        _, count = descriptions.export_template(RECORDS, template, STAGES)

        print("=" * 74)
        print("TEMPLATE")
        print("=" * 74)
        ws = load_workbook(template)["Descriptions"]
        print(f"  {len(RECORDS)} records -> {count} rows "
              f"(one per distinct target field)")
        print(f"  columns: {[c.value for c in ws[1]]}")

        # The human, being a human: wrong case, leading/trailing padding, a
        # non-breaking space, Vietnamese in NFD, and a row for a field that
        # does not exist.
        nbsp_text = "Số tham chiếu giao dịch"
        nfd_text = unicodedata.normalize("NFD", "Ngày mở")
        fill(template, {
            WIDE: ("  dwh_party ", "wide_field", nbsp_text),
            ("DWH_PARTY", "PLAIN_FIELD"): ("DWH_PARTY", "PLAIN_FIELD", nfd_text),
        }, extra_rows=[["DWH_GONE", "OLD_COL", "dwh", "", "", "Mồ côi"]])

        overlay = descriptions.load(template, STAGES)
        report = descriptions.merge_report(RECORDS, overlay)

        print()
        print("=" * 74)
        print("LOADED BACK")
        print("=" * 74)
        print(f"  entries read       : {len(overlay)}")
        print(f"  matched            : {report['matched']}/{report['fields']}")
        print(f"  orphaned           : {report['orphaned']} "
              f"{report['orphaned_keys']}")
        print(f"  undescribed        : {report['undescribed']} "
              f"{report['undescribed_keys']}")

        # the spreadsheet join
        book = tmp / "out.xlsx"
        write_workbook(RECORDS, str(book), STAGES, overlay=overlay)
        sheet = load_workbook(book).active
        headers = [c.value for c in sheet[3]]
        authored_col = headers.index(AUTHORED_COL) + 1
        written = [sheet.cell(r, authored_col).value
                   for r in range(4, sheet.max_row + 1)]
        fanned = sum(1 for v in written if v == nbsp_text.replace(" ", " "))

        plain = tmp / "plain.xlsx"
        write_workbook(RECORDS, str(plain), STAGES)
        plain_cols = load_workbook(plain).active.max_column

        print(f"\n  spreadsheet: {plain_cols} columns without an overlay, "
              f"{sheet.max_column} with")
        print(f"  the n->1 field's one meaning reached {fanned} record row(s)")

        # a file written for another layout
        other = tmp / "other.xlsx"
        descriptions.export_template(RECORDS, other, ["source", "warehouse"])
        try:
            descriptions.load(other, STAGES)
            refused = False
        except ValueError:
            refused = True

        # a blank authored cell is UNFILLED, not "deliberately empty"
        blank = tmp / "blank.xlsx"
        descriptions.export_template(RECORDS, blank, STAGES)
        blank_overlay = descriptions.load(blank, STAGES)

        checks = [
            ("one row per distinct target field, not per record",
             count == 3 and len(RECORDS) == 5),
            ("a key retyped with wrong case and padding still matches",
             report["matched"] == 2),
            ("a non-breaking space does not stop a cell being read",
             overlay.entries[("DWH_PARTY", "WIDE_FIELD")]
             == nbsp_text.replace(" ", " ")),
            ("Vietnamese in NFD matches the same text in NFC",
             overlay.entries[("DWH_PARTY", "PLAIN_FIELD")]
             == unicodedata.normalize("NFC", nfd_text)),
            ("a description for a field that no longer exists is ORPHANED",
             report["orphaned"] == 1
             and report["orphaned_keys"] == ["DWH_GONE.OLD_COL"]),
            ("an orphan is reported, never silently dropped",
             "DWH_GONE.OLD_COL" in report["orphaned_keys"]),
            ("a field with neither kind of description is counted",
             report["undescribed"] == 0),
            ("a document-stated description is carried into the template",
             any(r[DOCUMENT_COL] == "Ngày mở tài khoản"
                 for r in descriptions.field_rows(RECORDS))),

            # the two properties that must not break
            ("ONE description reaches ALL of a n->1 field's records",
             fanned == 3),
            ("the authored column exists only when an overlay is given",
             sheet.max_column == plain_cols + 1),
            ("the authored text never enters the records themselves",
             all(r.get("description") in (None, "Ngày mở tài khoản")
                 for r in RECORDS)),

            # refusals
            ("a file written against another layout is refused",
             refused),
            ("a blank authored cell means unfilled, not described",
             len(blank_overlay) == 0),
            # the template is made to be SENT, so it must not carry a home
            # directory. Every real path happens to contain "/Archive/", which
            # is why the fallback branch went unexercised until it was written
            # into a deliverable.
            ("the evidence column is shortened, never an absolute home path",
             all("/Users/" not in (r[descriptions.EVIDENCE_COL] or "")
                 for r in descriptions.field_rows(RECORDS))),
            ("a path outside any Archive folder is made home-relative",
             short_path(str(Path.home() / "Documents" / "x.xlsx"))
             == "~/Documents/x.xlsx"),
            ("a path outside home is left alone rather than mangled",
             short_path("/opt/shared/z.xlsx") == "/opt/shared/z.xlsx"),

            ("matching is by key, so re-sorting the sheet cannot misattribute",
             descriptions.field_key(*("  dwh_party ", "wide_field"))
             == descriptions.field_key("DWH_PARTY", "WIDE_FIELD")),
        ]

        print()
        print("=" * 74)
        print("ASSERTIONS")
        print("=" * 74)
        for label, ok in checks:
            print(f"  {'PASS' if ok else 'FAIL'}  {label}")
            failed += not ok
        print(f"\n  {len(checks) - failed}/{len(checks)} passed")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main_())
